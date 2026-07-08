from pathlib import Path


HEADERS = ["编号", "审计模块", "访谈对象", "访谈问题", "制度依据答案", "来源文件", "来源位置", "原文摘录", "置信度", "建议追问", "风险提示"]


def export_xlsx(path: Path, checklist: list[dict], missing: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "访谈清单"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center")
    for index, item in enumerate(checklist, start=1):
        sheet.append(
            [
                index,
                item.get("module"),
                item.get("interview_role"),
                item.get("interview_question"),
                item.get("expected_answer"),
                item.get("source_file"),
                item.get("source_location"),
                item.get("evidence_quote"),
                item.get("confidence"),
                item.get("follow_up_question"),
                item.get("risk_hint"),
            ]
        )
    for column in range(1, len(HEADERS) + 1):
        sheet.column_dimensions[chr(64 + column)].width = 18 if column < 4 else 32

    missing_sheet = workbook.create_sheet("制度缺失")
    missing_sheet.append(["模块", "问题", "风险", "建议访谈问题", "制度改进建议"])
    for item in missing:
        missing_sheet.append(
            [
                item.get("module"),
                item.get("issue"),
                item.get("risk"),
                item.get("suggested_interview_question"),
                item.get("suggested_policy_improvement"),
            ]
        )
    workbook.save(path)

