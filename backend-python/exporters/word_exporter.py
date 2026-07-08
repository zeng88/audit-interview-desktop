from pathlib import Path


def export_docx(path: Path, project: dict, checklist: list[dict], missing: list[dict]) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading("审计访谈清单", level=1)
    doc.add_paragraph(f"项目名称：{project.get('project_name', '')}")
    doc.add_paragraph(f"公司名称：{project.get('company_name', '')}")
    doc.add_paragraph(f"所属行业：{project.get('industry', '')}")
    doc.add_paragraph(f"审计目标：{project.get('audit_objective', '')}")
    doc.add_paragraph(f"审计期间：{project.get('audit_period', '')}")

    doc.add_heading("访谈清单", level=2)
    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    for idx, header in enumerate(["编号", "模块", "访谈对象", "访谈问题", "制度依据答案", "来源"]):
        table.rows[0].cells[idx].text = header
    for index, item in enumerate(checklist, start=1):
        row = table.add_row().cells
        row[0].text = str(index)
        row[1].text = item.get("module") or ""
        row[2].text = item.get("interview_role") or ""
        row[3].text = item.get("interview_question") or ""
        row[4].text = item.get("expected_answer") or ""
        row[5].text = f"{item.get('source_file') or ''} {item.get('source_location') or ''}"

    doc.add_heading("制度缺失清单", level=2)
    if missing:
        for item in missing:
            doc.add_paragraph(f"{item.get('module') or ''}：{item.get('issue') or ''}", style=None)
    else:
        doc.add_paragraph("未识别到制度缺失项。")

    doc.add_heading("生成说明", level=2)
    doc.add_paragraph("本文件由本地审计访谈清单生成工具生成，制度依据仅来自已上传制度片段。")
    doc.save(path)

