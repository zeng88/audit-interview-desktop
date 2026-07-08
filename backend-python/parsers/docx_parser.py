from pathlib import Path

from parsers.text_cleaner import clean_text


def parse_docx(path: str | Path) -> list[dict]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("缺少 python-docx 依赖，请执行 pip install -r requirements.txt") from exc

    doc = Document(str(path))
    blocks: list[str] = []
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            blocks.append(paragraph.text.strip())
    for table in doc.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if values:
                blocks.append(" | ".join(values))
    text = clean_text("\n".join(blocks))
    return [
        {
            "page_number": None,
            "sheet_name": None,
            "section_title": "Word 文档",
            "raw_text": text,
        }
    ] if text else []

