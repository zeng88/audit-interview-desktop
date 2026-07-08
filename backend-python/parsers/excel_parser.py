from pathlib import Path

from parsers.text_cleaner import clean_text


def parse_xlsx(path: str | Path) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("缺少 openpyxl 依赖，请执行 pip install -r requirements.txt") from exc

    workbook = load_workbook(str(path), data_only=True, read_only=True)
    pages: list[dict] = []
    for sheet in workbook.worksheets:
        lines: list[str] = []
        for row in sheet.iter_rows(values_only=True):
            values = [str(value).strip() for value in row if value is not None and str(value).strip()]
            if values:
                lines.append(" | ".join(values))
        text = clean_text("\n".join(lines))
        if text:
            pages.append(
                {
                    "page_number": None,
                    "sheet_name": sheet.title,
                    "section_title": sheet.title,
                    "raw_text": text,
                }
            )
    workbook.close()
    return pages

